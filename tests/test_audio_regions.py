"""The Phase 4 update for E2CL: one mixed file with several voices, crosstalk, music and effects.

Synthetic episodes (tests/audio_synth.py make_episode): two voices in different rooms, a music
intro and outro, a music bed under two sentences, an SFX hit, the client's stinger, crosstalk, an
echo by the other speaker and a turn with no pause. The clean ones must give no flags; real
mistakes planted next to all that must still be flagged. Not real E2CL audio: see the plan file.
"""

from __future__ import annotations

import json
import os
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import numpy as np  # noqa: E402
import pytest  # noqa: E402

from src.audio import regions as R  # noqa: E402
from src.audio.edits import analyze, check_audio, effective_settings  # noqa: E402
from src.models import SpokenWord  # noqa: E402
from tests import audio_synth as S  # noqa: E402
from tests.test_audio_edits import SETTINGS, write_wav  # noqa: E402

REF = [("stinger.wav", S.stinger())]


def check(ep_or_speech, settings=SETTINGS, **kw):
    sp = ep_or_speech
    kw.setdefault("references", REF)
    return check_audio(sp.samples, sp.sr, sp.words, settings, **kw)


def planted_episode(seed: int) -> tuple[S.Episode, S.Planter]:
    """A clean episode with three real mistakes: a hard cut mid-word in voice B's sentence 5 (no
    music near), sentence 4 said again after sentence 5 (a retake), a stutter in voice A's clear
    sentence 0. Planted from the end backwards, so the true times stay right."""
    ep = S.make_episode(seed)
    p = S.Planter(ep.speech)
    p.plant("abrupt_cut", S.cut_mid_word, ep.word_in(p.speech.words, 5, -3))
    sp = p.speech
    i4, j5 = ep.word_in(sp.words, 4, 0), ep.word_in(sp.words, 5, -1)
    k4 = max(i for i in range(len(sp.words)) if sp.words[i].start < ep.spans[4][1])
    sp.sentences = [(0, 0)] * 4 + [(i4, k4 + 1), (k4 + 1, j5 + 1)]
    p.plant("repeated_phrase", S.retake, 4, 5)
    p.plant("stutter", S.repeat_word, ep.word_in(p.speech.words, 0, 2))
    return ep, p


# --- clean episodes: zero flags -------------------------------------------------------------------

@pytest.mark.parametrize("seed", range(1, 7))
def test_clean_episode_with_voices_music_effects_and_crosstalk_gives_no_flags(seed):
    ep = S.make_episode(seed)
    issues, stats = check(ep)
    assert issues == [], [(i.timestamp, i.type, i.signals) for i in issues]
    kinds = {r["kind"] for r in stats["regions"]}
    assert {"music", "cue", "overlap"} <= kinds
    assert stats["lyrics_dropped"] == 4                               # the "la la la la" over the intro


@pytest.mark.parametrize("parts", [
    dict(intro=False, outro=False, bed=False, sfx=False, sting=False, overlap=True, echo=True),   # (a) two voices mixed
    dict(intro=False, outro=False, bed=True, sfx=False, sting=False, overlap=False, echo=False),  # (b) a bed
    dict(intro=True, outro=True, bed=False, sfx=False, sting=False, overlap=False, echo=False),  # (c) intro, outro
    dict(intro=False, outro=False, bed=False, sfx=True, sting=False, overlap=False, echo=False),  # (d) an SFX hit
    dict(intro=False, outro=False, bed=False, sfx=False, sting=False, overlap=False, echo=False),  # (e) speaker change
    dict(intro=False, outro=False, bed=False, sfx=False, sting=True, overlap=False, echo=False),  # (f) the stinger
])
def test_each_fixture_on_its_own_gives_no_flags(parts):
    ep = S.make_episode(11, lyrics=False, **parts)
    issues, _ = check(ep)
    assert issues == [], [(i.timestamp, i.type, i.signals) for i in issues]


def test_without_suppression_the_same_clean_episodes_are_flagged():
    """What the suppression is for: the old rules on the same clean audio."""
    flagged = 0
    for seed in range(1, 5):
        issues, stats = check(S.make_episode(seed), {**SETTINGS, "suppression": False})
        flagged += len(issues)
        assert stats["regions"] and stats["hidden"] == {}                  # regions found, nothing hidden
    assert flagged >= 8


# --- real mistakes next to all that: still flagged ----------------------------------------------------

def test_planted_mistakes_next_to_music_effects_and_voices_are_still_found():
    found, total, extra = {}, {}, 0
    for seed in range(1, 9):
        ep, p = planted_episode(seed)
        issues, _ = check(p.speech)
        for kind, t in p.truth:
            total[kind] = total.get(kind, 0) + 1
            found[kind] = found.get(kind, 0) + any(abs(i.timestamp - t) <= 1.5 and i.type == kind for i in issues)
        extra += sum(1 for i in issues if all(abs(i.timestamp - t) > 1.5 for _, t in p.truth))
    assert extra == 0
    assert found["stutter"] == total["stutter"]                        # in a clear section: all 8
    assert found["repeated_phrase"] >= total["repeated_phrase"] - 1      # 7-8 of 8
    assert found["abrupt_cut"] >= total["abrupt_cut"] - 2                # 6 of 8 (two sit at a speaker change)


def test_a_cut_far_from_any_cue_is_flagged_even_with_music_elsewhere():
    ep = S.make_episode(3)
    p = S.Planter(ep.speech)
    at = p.plant("abrupt_cut", S.cut_mid_word, ep.word_in(p.speech.words, 5, -3))
    issues, stats = check(p.speech)
    assert any(r["kind"] == "music" for r in stats["regions"])
    assert any(i.type == "abrupt_cut" and abs(i.timestamp - at) < 0.5 for i in issues)
    cut = next(i for i in issues if i.type == "abrupt_cut")
    assert "cut off mid-word" in cut.signals and cut.confidence >= 0.65  # each issue says what it rests on


# --- voices ----------------------------------------------------------------------------------------------

def test_a_word_repeated_by_the_other_speaker_is_not_a_stutter():
    """A says "...exactly." and B echoes "exactly": the same word twice, two voices."""
    rng = np.random.default_rng(5)
    sr = S.SR
    a = np.concatenate([S._syllable(rng, sr, 0.2, -20, S.VOICE_A) for _ in range(2)])
    b = np.concatenate([S._syllable(rng, sr, 0.2, -20, S.VOICE_B) for _ in range(2)])
    gap = S.room_tone(rng, int(0.15 * sr), -58)
    lead = S.make_speech(5, n_sentences=2)
    x = np.concatenate([lead.samples, a, gap, b, S.room_tone(rng, sr, -58)]).astype(np.float32)
    t0 = len(lead.samples) / sr
    words = lead.words + [SpokenWord("exactly.", t0, t0 + 0.4, 0.9), SpokenWord("exactly.", t0 + 0.55, t0 + 0.95, 0.9)]
    sp = S.Speech(x, words, [], sr)
    issues, stats = check(sp, references=())
    assert not [i for i in issues if i.type in ("stutter", "repeated_phrase")]
    assert stats["hidden"].get("repeat by a different voice (an echo, not a stutter)") == 1
    same = S.Speech(np.concatenate([lead.samples, a, gap, a, S.room_tone(rng, sr, -58)]).astype(np.float32), words, [], sr)
    issues, _ = check(same, references=())
    assert [i.type for i in issues] == ["stutter"]                     # the same voice twice: a stutter


def test_room_tone_alone_at_a_speaker_change_is_not_a_cut():
    """B is recorded in a noisier room (+8 dB) and answers with no pause."""
    ep = S.make_episode(7, intro=False, outro=False, bed=False, sfx=False, sting=False, overlap=False, echo=False)
    issues, _ = check(ep)
    assert issues == []
    issues, _ = check(ep, {**SETTINGS, "suppression": False})
    assert any("background noise changes" in i.signals for i in issues)   # what the old rule flagged


def test_speaker_changes_and_voice_distance():
    ep = S.make_episode(2, intro=False, outro=False, bed=False, sfx=False, sting=False, overlap=False, echo=False)
    s = effective_settings(SETTINGS)
    sig = analyze(ep.samples, ep.sr, s)
    f0 = R.pitch_track(sig.samples, sig.sr)
    found = R.speaker_changes(f0, ep.words, s)
    hits = sum(1 for t in ep.turns if any(abs(t - c) <= 1.0 for c in found))       # turns sit in pauses
    assert hits >= len(ep.turns) - 1 and len(found) <= len(ep.turns) + 2
    one_voice = S.make_speech(9, n_sentences=8)
    sig1 = analyze(one_voice.samples, one_voice.sr, s)
    assert R.speaker_changes(R.pitch_track(sig1.samples, sig1.sr), one_voice.words, s) == []


# --- music, effects, cues, cue sheets ---------------------------------------------------------------------

def test_regions_match_what_was_placed():
    ep = S.make_episode(4)
    _, stats = check(ep)
    regions = stats["regions"]

    def covered(a, b, kinds, share=0.9):
        """At least `share` of [a, b] lies in regions of these kinds."""
        inside = sum(max(0.0, min(b, r["end"]) - max(a, r["start"])) for r in regions if r["kind"] in kinds)
        return inside >= share * (b - a)

    for a, b, what in ep.intentional:
        if what in ("music intro", "music outro"):
            assert covered(a, b, ("music",)), what
        elif what == "music bed under speech":                         # found from its pauses: most of it
            assert covered(a, b, ("music",), share=0.5), what
        elif what == "stinger":
            assert any(r["kind"] == "cue" and abs(r["start"] - a) < 0.2 for r in regions)
        elif what == "crosstalk":
            assert covered(a, b, ("overlap",))
    speech = [r for r in regions if r["kind"] == "music" and any(r["start"] < w.start < r["end"] and w.probability > 0.5
                                                                   for w in ep.words)]
    assert all("bed" in r["why"] for r in speech)                    # music over words only where the bed is


def test_reference_matching_finds_the_stinger_and_nothing_else():
    rng = np.random.default_rng(1)
    sr = S.SR
    x = np.concatenate([S.room_tone(rng, 5 * sr, -58), S.stinger(), S.room_tone(rng, 4 * sr, -58),
                        S.music(3, 4.0, -18), S.stinger() * 0.5, S.room_tone(rng, 2 * sr, -58)]).astype(np.float32)
    found = R.cue_regions(x, sr, REF, {})
    assert [round(r.start, 1) for r in found] == [5.0, 14.6]
    assert R.cue_regions(S.music(4, 10.0, -18), sr, REF, {}) == []


def test_cue_sheet_wins_and_hides_only_its_own_span(tmp_path):
    ep = S.make_episode(3, intro=False, outro=False, bed=False, sfx=False, sting=False, overlap=False, echo=False)
    p = S.Planter(ep.speech)
    at = p.plant("abrupt_cut", S.cut_mid_word, ep.word_in(p.speech.words, 5, -3))
    sheet = tmp_path / "ep.cues.csv"
    sheet.write_text(f"start,end,type,note\n{at - 0.5:.2f},{at + 0.5:.2f},sfx,a scratch effect on purpose\n"
                     "0:01,0:02,banana,bad type\n", encoding="utf-8")
    regions, problems = R.read_cue_sheet(sheet)
    assert len(regions) == 1 and regions[0].source == "cue sheet" and len(problems) == 1
    issues, stats = check(p.speech, cue_sheet=regions)
    assert not [i for i in issues if i.type == "abrupt_cut"]
    assert any(r["source"] == "cue sheet" and r["suppressed"] == 1 for r in stats["regions"])
    far = [R.Region(at + 5, at + 6, "sfx", "cue sheet", "elsewhere")]
    issues, _ = check(p.speech, cue_sheet=far)
    assert any(i.type == "abrupt_cut" for i in issues)               # a cue somewhere else hides nothing


def test_cue_sheets_are_attached_next_to_the_episode(tmp_path):
    ep, csv_file = tmp_path / "Episode 12.mp3", tmp_path / "cues.csv"
    ep.write_bytes(b"x")
    csv_file.write_text("start,end,type,note\n0,8,music,intro\n", encoding="utf-8")
    assert R.match_cue_sheets([ep], [csv_file]) == {ep: csv_file}
    where, problems = R.attach_cue_sheet(ep, csv_file, tmp_path / "fallback")
    assert where == tmp_path / "Episode 12.cues.csv" and problems == []
    assert R.find_cue_sheet(ep) == where
    named = tmp_path / "Episode 12 cues.csv"
    named.write_text("start,end,type\n1,2,sfx\n", encoding="utf-8")
    assert R.match_cue_sheets([ep, tmp_path / "other.mp3"], [named, csv_file]) == {ep: named}
    bad = tmp_path / "bad.csv"
    bad.write_text("a,b\n1,2\n", encoding="utf-8")
    with pytest.raises(ValueError):
        R.attach_cue_sheet(ep, bad)


def test_guidelines_decide_what_counts_as_intentional():
    ep = S.make_episode(4, intro=False, outro=False, sfx=True, sting=False, overlap=False, echo=False)
    _, allowed = check(ep, guide={"music_allowed": ["under_speech"], "sfx_allowed": ["emphasis"]})
    _, strict = check(ep, guide={"music_allowed": ["intro", "outro"], "sfx_allowed": []})
    assert any(r["kind"] == "music" and "bed" in r["why"] for r in allowed["regions"])
    assert not any(r["kind"] == "music" and "bed" in r["why"] for r in strict["regions"])
    assert any(r["kind"] == "unexpected" and "bed" in r["why"] for r in strict["regions"])
    margins = [R.guideline_settings({"sensitivity": s})["margin_sec"] for s in ("strict", "normal", "lenient")]
    assert margins == [0.15, 0.3, 0.6]


def test_required_cues_check_is_off_by_default_and_reports_a_missing_one():
    ep = S.make_episode(5, sting=False)
    guide = {"required_cues": {"enabled": True, "cues": [{"reference": "stinger.wav", "at": "start", "within_sec": 30}]}}
    assert not [i for i in check(ep)[0] if i.type == "missing_cue"]
    issues, _ = check(ep, guide=guide)
    assert [i.type for i in issues] == ["missing_cue"] and "stinger.wav" in issues[0].message
    with_cue = S.make_episode(5)                                       # the stinger at ~17 s: within 30 s
    assert not [i for i in check(with_cue, guide=guide)[0] if i.type == "missing_cue"]


# --- engine, report, Whisper segments -----------------------------------------------------------------------

def test_report_lists_the_ignored_regions_and_the_overlap_note(tmp_path, monkeypatch):
    from docx import Document

    from src.config import load_config
    from src.engine import run_qc, write_report

    ep, p = planted_episode(2)
    path = write_wav(tmp_path / "E2CL ep.wav", p.speech.samples, p.speech.sr)
    monkeypatch.setattr("src.engine.get_transcript", lambda *a, **k: p.speech.words)
    config = load_config()
    config = {**config, "report": {**config["report"], "output_dir": str(tmp_path / "out")}}
    result = run_qc(path, "E2CL", config=config)
    assert any(line.startswith("Regions ignored by the bot:") for line in result.checks)
    written = write_report(result, formats=["md", "docx"])
    md = next(x for x in written if x.suffix == ".md").read_text(encoding="utf-8")
    assert "## Regions ignored by the bot" in md and "isn't checked for stutters" in md
    assert "| Overlapping speech |" in md and "| Music |" in md and "*Signals:" in md
    doc = Document(str(next(x for x in written if x.suffix == ".docx")))
    texts = [par.text for par in doc.paragraphs]
    assert "Regions ignored by the bot" in texts
    headers = [[c.text for c in t.rows[0].cells] for t in doc.tables]
    assert ["Start", "End", "What", "Why", "Flags hidden"] in headers
    csv_text = next(x for x in written if x.name == "issues.csv").read_text(encoding="utf-8-sig")
    assert csv_text.splitlines()[0].endswith(",signals")


def test_whisper_segments_are_kept_and_old_caches_are_redone(tmp_path):
    from src.audio import transcribe

    audio = tmp_path / "a.wav"
    write_wav(audio, np.zeros(16000, np.float32), 16000)
    calls = []

    def fake(wav, ws, prompt):
        calls.append(1)
        out = transcribe.Transcript([SpokenWord("hi", 0.1, 0.3, 0.9)], [(0.0, 1.0, 0.85)])
        return out

    machine = {"whisper_model_size": "base", "cpu_threads": 2}
    monkey_ffmpeg = transcribe.find_ffmpeg
    words = transcribe.get_transcript(audio, machine, {}, [], tmp_path / "cache", transcriber=fake)
    assert words.segments == [(0.0, 1.0, 0.85)] and len(calls) == 1
    cached = next((tmp_path / "cache").glob("*.json"))
    data = json.loads(cached.read_text(encoding="utf-8"))
    del data["segments"]                                         # a transcript cached before segments were kept
    cached.write_text(json.dumps(data), encoding="utf-8")
    again = transcribe.get_transcript(audio, machine, {}, [], tmp_path / "cache", transcriber=fake)
    assert again.segments is None and len(calls) == 1             # fine for the caption check
    again = transcribe.get_transcript(audio, machine, {}, [], tmp_path / "cache", transcriber=fake, need_segments=True)
    assert again.segments == [(0.0, 1.0, 0.85)] and len(calls) == 2
    assert monkey_ffmpeg is transcribe.find_ffmpeg


def test_music_segments_from_whisper_count_and_lyrics_are_dropped():
    """A stretch Whisper says has no speech (while it isn't silent) is music, and its unsure words go."""
    rng = np.random.default_rng(2)
    sr = S.SR
    lead = S.make_speech(8, n_sentences=2)
    pad = S.music(9, 5.0, -18) + S.room_tone(rng, 5 * sr, -58)
    x = np.concatenate([lead.samples, pad]).astype(np.float32)
    t0 = len(lead.samples) / sr
    ghost = [SpokenWord("oh", t0 + 1 + k * 0.8, t0 + 1.3 + k * 0.8, 0.4) for k in range(4)]
    issues, stats = check_audio(x, sr, lead.words + ghost, SETTINGS, segments=[(t0, t0 + 5, 0.9)])
    assert issues == [] and stats["lyrics_dropped"] == 4
    assert any(r["kind"] == "music" and r["start"] <= t0 + 0.5 for r in stats["regions"])


def test_template_and_dialog_carry_the_audio_guidelines(tmp_path):
    from openpyxl import load_workbook

    from src.clients.importers import import_file, make_template
    from src.clients.model import Finding, build_client

    path = make_template(tmp_path / "t.xlsx")
    assert import_file(path) and not [f for f in import_file(path) if "audio_guidelines" in str(f.value)]
    wb = load_workbook(path)
    ws = wb["Audio guidelines"]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "music_allowed":
            row[1].value = "intro, outro, ad breaks"
        if row[0].value == "sensitivity":
            row[1].value = "lenient"
    wb.save(path)
    found = import_file(path) + [Finding("code", "E9POD"), Finding("name", "Pod")]
    client = build_client(found).client
    assert client["audio_guidelines"] == {"music_allowed": ["intro", "outro", "ad_breaks"], "sensitivity": "lenient"}
    assert "audio_guidelines" not in client["overrides"]
