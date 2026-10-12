"""The ignore list (ignore.yaml): reading rules, matching flags, the lines the report offers, and
the engine hiding flags and saying how many."""

from pathlib import Path

import pytest
from docx import Document

from src import ignore
from src.config import load_config
from src.engine import QCOptions, build_docx_report, run_qc
from src.models import Issue

FIXTURES = Path(__file__).resolve().parent / "fixtures"
SRT = FIXTURES / "sample_mistakes.srt"


def issue(kind="spelling", text="Meet Gherabel tonight", t=10.0, cid=4):
    return Issue(cid, t, kind, "medium", text, "msg", "", 0.8)


def rules_from(tmp_path, body: str) -> list[ignore.Rule]:
    path = tmp_path / "ignore.yaml"
    path.write_text(body, encoding="utf-8")
    return ignore.load_rules(path)


def test_no_file_or_an_empty_list_ignores_nothing(tmp_path):
    assert ignore.load_rules(tmp_path / "missing.yaml") == []
    assert rules_from(tmp_path, "ignore:\n") == []
    assert rules_from(tmp_path, "# only comments\n") == []


def test_project_ignore_file_starts_empty():
    config = load_config()
    assert ignore.ignore_file(config).is_file() and ignore.load_rules(ignore.ignore_file(config)) == []


def test_broken_rules_are_skipped_not_fatal(tmp_path, caplog):
    rules = rules_from(tmp_path, """ignore:
  - {type: spelling, text: Gherabel}
  - {type: spelling, colour: red}
  - {note: "matches everything"}
  - {type: color_cast, at: "soon"}
  - just a string
""")
    assert len(rules) == 1 and rules[0].text == "Gherabel"
    assert sum(r.getMessage().endswith(("skipped", "time)")) for r in caplog.records) == 4
    assert rules_from(tmp_path, "ignore: [unclosed\n") == []                  # unreadable YAML: nothing hidden


def test_every_given_field_must_match(tmp_path):
    rule, = rules_from(tmp_path, "ignore:\n  - {client: e2ab, type: spelling, text: GHERABEL}\n")
    assert rule.matches(issue(text="meet   gherabel tonight"), "E2AB", None)        # any case and spacing
    assert not rule.matches(issue(), "E8D", None)                                     # another client
    assert not rule.matches(issue(), None, None)
    assert not rule.matches(issue(kind="capitalization"), "E2AB", None)               # another type
    assert not rule.matches(issue(text="Meet Gabriel tonight"), "E2AB", None)         # other words


def test_video_and_time_rules(tmp_path):
    rule, = rules_from(tmp_path, 'ignore:\n  - {video: "Willy & Mitzie.MP4", type: color_cast, at: "2:04"}\n')
    flag = issue(kind="color_cast", text="Shot 7", t=125.5, cid=-1007)
    assert rule.matches(flag, "E3AC", Path("D:/QC/willy & mitzie.mp4"))
    assert not rule.matches(issue(kind="color_cast", t=130.0, cid=-1007), "E3AC", Path("Willy & Mitzie.mp4"))
    assert not rule.matches(flag, "E3AC", Path("other.mp4"))
    assert not rule.matches(flag, "E3AC", None)


def test_apply_splits_kept_and_hidden(tmp_path):
    rules = rules_from(tmp_path, "ignore:\n  - {type: spelling, text: Gherabel}\n")
    a, b = issue(), issue(text="Gabriel arrives")
    assert ignore.apply([a, b], rules, None, None) == ([b], [a])


@pytest.mark.parametrize("flag", [
    issue(text='He said "stop: now" # really'),                            # quotes, colon, hash
    issue(kind="punctuation", text="it's 5 o'clock\nagain"),               # apostrophes, a line break
    issue(kind="color_cast", text="Shot 3", t=64.25, cid=-1003),           # a color flag: video and time
    issue(kind="uncaptioned_speech", text="", t=12.0, cid=-2),             # no caption at all
])
def test_report_lines_hide_exactly_their_flag(tmp_path, flag):
    video = Path("D:/VidÃ©os/Willy & Mitzie: final.mp4")
    line = ignore.rule_line(flag, "E3AC", video)
    rules = rules_from(tmp_path, "ignore:\n  " + line + "\n")
    assert len(rules) == 1, line
    assert rules[0].matches(flag, "E3AC", video)
    assert not rules[0].matches(flag, "E8D", video)                        # tied to the client


def test_engine_hides_ignored_flags_and_says_so(tmp_path):
    config = load_config()
    before = run_qc(None, "E8D", options=QCOptions(srt_path=SRT, use_llm=False), config=config)
    target = before[0]
    path = tmp_path / "my ignore.yaml"
    path.write_text("ignore:\n  " + ignore.rule_line(target, "E8D", None) + "\n", encoding="utf-8")
    config["paths"]["ignore_file"] = str(path)                             # an absolute path works too
    after = run_qc(None, "E8D", options=QCOptions(srt_path=SRT, use_llm=False), config=config)
    assert target not in after and target in after.hidden
    assert len(after) + len(after.hidden) == len(before)
    assert any(line.startswith("Ignore list: ") and "hidden" in line for line in after.checks)
    assert not any(line.startswith("Ignore list") for line in before.checks)


def test_report_offers_a_line_per_flag(tmp_path):
    result = run_qc(None, "E8D", options=QCOptions(srt_path=SRT, use_llm=False))
    doc = Document(str(build_docx_report(result, output_dir=tmp_path)))
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert headings[-1] == "Hide a flag next time"
    lines = doc.tables[-1]
    assert len(lines.rows) == 1 + len(result)
    assert lines.rows[1].cells[1].text.startswith("- {client: E8D, type: ")


# --- unknown words: ignore just the word, or add it to the glossary ---------------------------

def spelling_flag(message, text="We met Mark McGwire today"):
    return Issue(3, 5.0, "spelling", "low", text, message, "", 0.5)


def test_the_flagged_word_comes_from_the_message():
    from src.checks.spelling import flagged_word

    assert flagged_word(spelling_flag('Unknown word "McGwire": a misspelling, or a name to add to glossary.txt?')) == \
        "McGwire"
    assert flagged_word(spelling_flag('Possible misspelling: "hourse".')) == "hourse"
    assert flagged_word(spelling_flag('"McGuire" should be spelled "McGwire" (glossary).')) is None
    assert flagged_word(issue(kind="capitalization")) is None


def test_an_unknown_word_rule_covers_every_caption_with_it(tmp_path):
    flag = spelling_flag('Unknown word "McGwire": a misspelling, or a name to add to glossary.txt?')
    line = ignore.rule_line(flag, "E7SS", None)
    assert line == '- {client: E7SS, type: spelling, text: "McGwire"}'
    rule, = rules_from(tmp_path, "ignore:\n  " + line + "\n")
    other = spelling_flag(flag.message, text="McGwire hit 70 home runs")
    assert rule.matches(flag, "E7SS", None) and rule.matches(other, "E7SS", None)
    assert ignore.glossary_line(flag) == "McGwire" and ignore.glossary_line(issue(kind="capitalization")) is None


def test_report_offers_glossary_lines_for_unknown_words(tmp_path):
    result = run_qc(None, "E8D", options=QCOptions(srt_path=SRT, use_llm=False))
    doc = Document(str(build_docx_report(result, output_dir=tmp_path)))
    lines = doc.tables[-1]
    assert [c.text for c in lines.rows[0].cells] == ["Time", "Line for ignore.yaml", "Or add to the glossary"]
    from src.checks.spelling import flagged_word

    for row, flag in zip(lines.rows[1:], result):
        assert row.cells[2].text == (flagged_word(flag) or "")
