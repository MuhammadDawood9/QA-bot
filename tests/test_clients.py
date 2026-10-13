"""Phase 2: adding and editing clients. Importing the template, the QC review workbook, CSV and
PDF; reviewing; saving with validation, duplicate check, backups and the changelog."""

from __future__ import annotations

import os
import shutil
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest  # noqa: E402
import yaml  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

from src.clients.importers import (TEMPLATE_SHEETS, detect_from_text, import_file, make_template,  # noqa: E402
                                   read_with_claude)
from src.clients.model import (ClientError, Finding, build_client, describe_changes, finding_text,  # noqa: E402
                               findings_from_client, merge_findings, parse_area, parse_finding_text)
from src.clients.store import ClientStore, write_color_clients  # noqa: E402
from src.config import color_settings, load_client, load_config  # noqa: E402
from src.engine import run_qc  # noqa: E402
from src.utils.safe_io import read_changelog  # noqa: E402
from tests.pdf_helper import make_text_pdf  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
FIXTURES = Path(__file__).resolve().parent / "fixtures"
REAL_XLSX = ROOT / "data" / "QC_Review_Analysis_Jun-Oct_2026.xlsx"


@pytest.fixture
def data(tmp_path) -> dict:
    """A data folder like the real one: config.yaml, color.yaml, clients/E8D.yaml, glossary."""
    for name in ("config.yaml", "color.yaml", "rules.yaml", "glossary.txt", "extra_words.txt", "ignore.yaml"):
        shutil.copy(ROOT / name, tmp_path / name)
    (tmp_path / "clients").mkdir()
    shutil.copy(ROOT / "clients" / "E8D.yaml", tmp_path / "clients" / "E8D.yaml")
    return load_config(tmp_path / "config.yaml")


def filled_template(path: Path, code="E9ZZ") -> Path:
    make_template(path)
    wb = load_workbook(path)
    info = wb["Info"]
    info["B2"], info["B3"], info["B4"], info["B5"], info["B6"] = code, "Zed podcasts", "en", 3, "bottom"
    wb["Rules"].append(["Never put a full stop at the end of a caption", "Don't", ""])
    wb["Glossary"].append(["Danja", "no", ""])
    wb["Glossary"].append(["YouTube", "yes", ""])
    wb["Exceptions"].append(["spelling", "Gherabel", "", "", "guest's name"])
    wb["Exceptions"].append(["", "", "", "", "Long titles are fine"])
    wb["Examples"].append(["aging", "Aging", "capitalization", ""])
    wb["Checks"].append(["capitalization", "allow_stylized_lowercase", "yes", ""])
    wb["Checks"].append(["color", "skip", "saturation, color_cast", ""])
    wb["Checks"].append(["color", "cast_strength", 9, ""])
    wb["Checks"].append(["spelling", "enabled", "no", ""])
    wb.save(path)
    return path


# --- template ------------------------------------------------------------------------------------

def test_template_has_every_sheet_and_skips_its_example_rows(tmp_path):
    path = make_template(tmp_path / "client_template.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == list(TEMPLATE_SHEETS)
    found = import_file(path)
    assert {f.kind for f in found} <= {"language", "min_words", "caption_area"}   # only the Info defaults
    assert not [f for f in found if "e.g." in finding_text(f)]


def test_filled_template_becomes_a_client_file(tmp_path, data):
    found = import_file(filled_template(tmp_path / "t.xlsx"))
    by_kind = {}
    for f in found:
        by_kind.setdefault(f.kind, []).append(f.value)
    assert by_kind["code"] == ["E9ZZ"] and by_kind["name"] == ["Zed podcasts"] and by_kind["min_words"] == [3]
    assert by_kind["caption_area"] == [parse_area("bottom")]
    assert by_kind["forced"] == ["YouTube"] and by_kind["glossary"] == ["Danja"]
    assert by_kind["exception"] == [{"type": "spelling", "text": "Gherabel", "note": "guest's name"}]
    assert by_kind["note"] == ["Long titles are fine"]
    assert by_kind["color_skip"] == ["saturation", "color_cast"]

    store = ClientStore(data)
    changes = store.save(build_client(found), "add", sources=[tmp_path / "t.xlsx"], user="tester")
    assert "glossary" in changes and "imported t.xlsx" in changes
    saved = yaml.safe_load((Path(data["_root"]) / "clients" / "E9ZZ.yaml").read_text(encoding="utf-8"))
    assert saved["glossary"] == ["Danja", "=YouTube"]
    assert saved["overrides"] == {"capitalization": {"allow_stylized_lowercase": True},
                                  "spelling": {"enabled": False}, "too_short": {"min_words": 3}}
    assert saved["ocr"]["region"] == parse_area("bottom")
    assert saved["rules"] == ["Don't: Never put a full stop at the end of a caption"]
    # color goes to color.yaml, and the engine's color settings see it
    config = load_config(Path(data["_root"]) / "config.yaml")
    settings = color_settings(config, load_client(config, "9ZZ"))
    assert settings["skip"] == ["saturation", "color_cast"] and settings["cast_strength"] == 9
    assert settings["clipped_white"] == 0.04                               # untouched defaults stay
    text = (Path(data["_root"]) / "color.yaml").read_text(encoding="utf-8")
    assert text.startswith("# Color grading checks") and "overexposed_high_factor: 2.5" in text
    assert "Long titles are fine" in store.notes_file("E9ZZ").read_text(encoding="utf-8")
    assert list((Path(data["_root"]) / "clients" / "E9ZZ" / "sources").glob("*_t.xlsx"))
    [row] = read_changelog(Path(data["_root"]) / "changelog.csv")
    assert row["user"] == "tester" and row["client"] == "E9ZZ" and row["action"] == "add" and row["computer"]


def test_duplicate_codes_are_blocked_in_any_spelling(data):
    store = ClientStore(data)
    for code in ("E8D", "8d", " e8D "):
        with pytest.raises(ClientError, match="already a client E8D"):
            store.save(build_client([Finding("code", code), Finding("name", "x")]), "add")


def test_edit_backs_up_and_logs_what_changed(data):
    store = ClientStore(data)
    current = store.load("E8D")
    findings = findings_from_client(current, store.color_block("E8D"))
    findings.append(Finding("forced", "Danja"))
    findings = [f if f.kind != "name" else Finding("name", "Danja promos") for f in findings]
    changes = store.save(build_client(findings, base=current), "edit", user="editor1")
    assert "glossary: +=Danja" in changes and "name: " in changes
    backups = list((Path(data["_root"]) / "backups" / "clients" / "E8D").glob("E8D_*_editor1.yaml"))
    assert len(backups) == 1 and "Danja promo/credit videos" in backups[0].read_text(encoding="utf-8")
    saved = store.load("E8D")
    assert saved["glossary"] == ["=Danja"] and saved["project"] == "Danja promo/credit videos"   # unmanaged key kept
    assert saved["overrides"] == {"capitalization": {"allow_stylized_lowercase": True}}
    assert read_changelog(store.changelog)[-1]["user"] == "editor1"
    # saving the same thing again changes nothing and logs nothing
    assert store.save(build_client(findings_from_client(saved), base=saved), "edit") == ""
    assert len(read_changelog(store.changelog)) == 1


def test_validation_messages():
    ok = [Finding("code", "E5X"), Finding("name", "x")]
    for bad, message in [
        ([Finding("code", "?"), Finding("name", "x")], "Client code"),
        ([Finding("code", "E5X")], "Display name: required"),
        (ok + [Finding("min_words", 0)], "Minimum words"),
        (ok + [Finding("language", "klingon")], "Language"),
        (ok + [Finding("caption_area", {"left": 0.5, "top": 0.5, "width": 0.7, "height": 0.2})], "past the edge"),
        (ok + [Finding("exception", {"note": "only a note"})], "every flag"),
        (ok + [Finding("color_skip", "purple")], "Color type"),
        (ok + [Finding("check", {"setting": "nonsense.x", "value": 1})], "check name"),
        (ok + [Finding("name", "y")], "ticked twice"),
    ]:
        with pytest.raises(ClientError, match=message):
            build_client(bad)
    assert build_client(ok + [Finding("min_words", 0, keep=False)]).client["overrides"] == {}


def test_edited_text_round_trips():
    for f in [Finding("example", {"wrong": "aging", "correct": "Aging", "category": "capitalization"}),
              Finding("check", {"setting": "too_short.phrase_break.fragment_max_words", "value": 3}),
              Finding("color", {"setting": "skin_hue_range", "value": [15, 75]}),
              Finding("exception", {"type": "spelling", "text": "Gherabel"}),
              Finding("caption_area", parse_area("0.1, 0.7, 0.8, 0.25")),
              Finding("min_words", 3)]:
        assert parse_finding_text(f.kind, finding_text(f)) == f.value
    assert parse_finding_text("check", "spelling.enabled = no") == {"setting": "spelling.enabled", "value": False}
    with pytest.raises(ClientError):
        parse_finding_text("example", "no arrow here")


def test_merge_keeps_one_of_each_single_value():
    merged = merge_findings([Finding("min_words", 3, "pdf")], [Finding("min_words", 2, "typed")],
                            [Finding("glossary", "Danja"), Finding("glossary", "danja")])
    assert [(f.value, f.keep) for f in merged if f.kind == "min_words"] == [(3, True), (2, False)]
    assert len([f for f in merged if f.kind == "glossary"]) == 1


def test_color_clients_block_rewritten_without_touching_the_defaults(tmp_path):
    path = tmp_path / "color.yaml"
    shutil.copy(ROOT / "color.yaml", path)
    before = path.read_text(encoding="utf-8")
    write_color_clients(path, {"E3AC": {"skip": ["saturation"]}})
    after = path.read_text(encoding="utf-8")
    assert after.split("\nclients:")[0] == before.split("\nclients:")[0]
    assert yaml.safe_load(after)["clients"] == {"E3AC": {"skip": ["saturation"]}}
    write_color_clients(path, {})
    assert yaml.safe_load(path.read_text(encoding="utf-8"))["clients"] == {}


def test_describe_changes():
    line = describe_changes({"glossary": ["a"], "overrides": {"too_short": {"min_words": 2}}},
                            {"glossary": ["a", "b"], "overrides": {"too_short": {"min_words": 3}}})
    assert line == "glossary: +b; overrides.too_short.min_words: 2 -> 3"


# --- the QC review workbook and CSV ----------------------------------------------------------------

def analysis_workbook(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Client rules"
    ws.append(["client_code", "project (from task name)", "client-specific rule"])
    ws.append(["2CS", "Market podcast clips", "All lowercase style allowed; minimum 3 words per caption"])
    ws.append(["E8D", "Danja", "other client"])
    ws = wb.create_sheet("Caption mistakes (2)")
    ws.append(["task_name", "client", "video_timestamp", "editor_wrote", "should_be", "category", "comment_text"])
    ws.append(["#2CS Episode 1", "", "00:01", "Warren buffet", "Warren Buffett", "spelling", "name"])
    ws.append(["#E8D reel", "E8D", "00:02", "x", "y", "spelling", ""])
    ws = wb.create_sheet("Accepted exceptions")
    ws.append(["client", "what you accepted"])
    ws.append(["E2CS", "0:12 long title on purpose"])
    wb.save(path)
    return path


def test_analysis_workbook_takes_one_clients_rows_and_normalizes_codes(tmp_path):
    found = import_file(analysis_workbook(tmp_path / "QC_Review_Analysis.xlsx"), "e2cs")
    kinds = [(f.kind, f.value, f.keep) for f in found]
    assert ("code", "E2CS", True) in kinds and ("name", "Market podcast clips", True) in kinds
    assert ("min_words", 3, True) in kinds
    assert ("check", {"setting": "capitalization.allow_stylized_lowercase", "value": True}, True) in kinds
    assert ("forced", "Warren Buffett", False) in kinds                     # a candidate: shown, not ticked
    assert any(k == "example" and v["wrong"] == "Warren buffet" for k, v, _ in kinds)
    assert ("note", "Accepted: 0:12 long title on purpose", True) in kinds
    assert not any(isinstance(v, dict) and v.get("wrong") == "x" for _, v, _ in kinds)       # E8D's row
    missing = import_file(tmp_path / "QC_Review_Analysis.xlsx", "E7SS")
    assert missing[0].kind == "note" and "Clients in it: E2CS, E8D" in missing[0].value


@pytest.mark.skipif(not REAL_XLSX.is_file(), reason="the real QC review workbook isn't on this machine")
def test_real_analysis_workbook():
    found = import_file(REAL_XLSX, "8D")
    assert ("code", "E8D") in [(f.kind, f.value) for f in found]
    assert any(f.kind == "rule" for f in found) and any(f.kind == "example" for f in found)
    assert any(f.kind == "check" and f.value["setting"] == "capitalization.allow_stylized_lowercase" for f in found)


def test_csv_from_a_template_sheet_and_unknown_csv(tmp_path):
    glossary = tmp_path / "glossary.csv"
    glossary.write_text("Word,Forced spelling (yes/no),Note\nDanja,no,\nYouTube,yes,\n", encoding="utf-8-sig")
    assert [(f.kind, f.value) for f in import_file(glossary)] == [("glossary", "Danja"), ("forced", "YouTube")]
    other = tmp_path / "other.csv"
    other.write_text("a,b\nhello,world\n", encoding="utf-8")
    assert [f.kind for f in import_file(other)] == ["note", "note"]


def test_unsupported_and_damaged_files_give_a_message(tmp_path):
    (tmp_path / "x.docx").write_text("x", encoding="utf-8")
    with pytest.raises(ClientError, match="only .xlsx, .csv and .pdf"):
        import_file(tmp_path / "x.docx")
    (tmp_path / "broken.xlsx").write_bytes(b"not a zip")
    with pytest.raises(ClientError, match="couldn't be read"):
        import_file(tmp_path / "broken.xlsx")
    (tmp_path / "broken.pdf").write_bytes(b"%PDF-1.4 nothing")
    with pytest.raises(ClientError):
        import_file(tmp_path / "broken.pdf")


# --- PDF ---------------------------------------------------------------------------------------

GUIDE = ["Caption style guide", "Captions sit at the bottom of the frame.", "Minimum of 3 words per caption.",
         "Names:", "Danja, Mark McGwire, YouTube", "", "Never animate the word the on its own.",
         "Always spell it as \"YouTube\".", "Language: English", "Logo safe zone is 10%."]


def test_pdf_detects_rules_and_keeps_the_rest_as_notes(tmp_path):
    found = import_file(make_text_pdf(tmp_path / "guide.pdf", GUIDE))
    got = {(f.kind, str(f.value)) for f in found if f.keep}
    assert ("min_words", "3") in got and ("language", "en") in got
    assert ("caption_area", str(parse_area("bottom"))) in got
    assert {("glossary", "Danja"), ("glossary", "Mark McGwire"), ("forced", "YouTube")} <= got
    assert ("rule", "Don't: Never animate the word the on its own.") in got
    assert ("note", "Logo safe zone is 10%.") in got and ("note", "Caption style guide") in got
    assert all(f.source.startswith("guide.pdf, line ") for f in found)


@pytest.mark.parametrize("line, words", [("At least two words on every caption", 2),
                                         ("2-4 words per caption", 2), ("3 words minimum", 3),
                                         ("No single-word captions", 2)])
def test_min_words_phrasings(line, words):
    assert [f.value for f in detect_from_text(line, "x") if f.kind == "min_words"] == [words]


def test_read_with_claude_uses_a_fake_client():
    class Fake:
        def complete_json(self, system, user, schema):
            return ('{"min_words": 3, "caption_position": "top", "language": "Spanish", "glossary": ["Danja"], '
                    '"forced_spellings": ["YouTube"], "rules": ["No emojis"]}')

    found = read_with_claude("guide text", "guide.pdf", {}, client=Fake())
    assert [(f.kind, f.value) for f in found] == [
        ("min_words", 3), ("caption_area", parse_area("top")), ("language", "es"), ("glossary", "Danja"),
        ("forced", "YouTube"), ("rule", "No emojis")]


def test_claude_button_is_off_by_default(monkeypatch):
    from src.clients.importers import claude_allowed

    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-test")
    assert not claude_allowed(load_config())
    assert claude_allowed({"client_import": {"allow_claude": True}})
    monkeypatch.delenv("ANTHROPIC_API_KEY")
    assert not claude_allowed({"client_import": {"allow_claude": True}})


# --- the engine uses what's saved -------------------------------------------------------------------

def test_client_ignore_rules_and_language_reach_the_engine(data):
    store = ClientStore(data)
    store.save(build_client([Finding("code", "E5X"), Finding("name", "x"),
                             Finding("exception", {"type": "spelling", "text": "Gherabel"})]), "add")
    config = load_config(Path(data["_root"]) / "config.yaml")
    srt = Path(data["_root"]) / "g.srt"
    srt.write_text("1\n00:00:01,000 --> 00:00:03,000\nWe met Gherabel today\n", encoding="utf-8")
    plain = run_qc(None, None, options=__import__("src.engine", fromlist=["QCOptions"]).QCOptions(srt_path=srt),
                   config=config)
    assert any(i.type == "spelling" for i in plain)
    hidden = run_qc(None, "E5X", options=__import__("src.engine", fromlist=["QCOptions"]).QCOptions(srt_path=srt),
                    config=config)
    assert not any(i.type == "spelling" for i in hidden) and len(hidden.hidden) == 1
    from src.engine import build_context

    assert build_context(config, {"client": "X", "language": "es"}).settings["spelling"]["language"] == "es"


# --- the dialogs (headless) ---------------------------------------------------------------------------

@pytest.fixture(scope="module")
def qapp():
    from src.app import main as app_main

    return app_main.make_application(["test"])


def test_review_dialog_reads_back_edits_and_ticks(qapp):
    from PySide6.QtCore import Qt

    from src.app.client_dialogs import ReviewDialog

    saved = []
    dialog = ReviewDialog([Finding("code", "E5X"), Finding("name", "x"), Finding("glossary", "Danja"),
                           Finding("min_words", 2)], lambda f: saved.append(f) or None)
    gloss = dialog.groups["glossary"].child(0)
    gloss.setText(0, "Danja Studio")
    dialog.groups["min_words"].child(0).setCheckState(0, Qt.CheckState.Unchecked)
    dialog.save()
    [findings] = saved
    assert ("glossary", "Danja Studio", True) in [(f.kind, f.value, f.keep) for f in findings]
    assert ("min_words", 2, False) in [(f.kind, f.value, f.keep) for f in findings]


def test_client_dialog_add_saves_through_the_store(qapp, data, tmp_path):
    from src.app.client_dialogs import ClientDialog

    store = ClientStore(data)
    dialog = ClientDialog(store, data, None)
    dialog.code.setText("7zz")
    dialog.name.setText("Seven")
    dialog.files.add_paths([make_text_pdf(tmp_path / "g.pdf", GUIDE), tmp_path / "nope.docx"])
    assert dialog.files.count() == 1
    from src.clients.importers import import_file as imp

    findings = merge_findings(dialog.typed_findings()[:2], imp(dialog.files.paths()[0]), dialog.typed_findings()[2:])
    assert dialog.save_findings(findings) is None and dialog.saved_code == "E7ZZ"
    assert store.load("E7ZZ")["overrides"]["too_short"]["min_words"] == 3          # the PDF's, not the default 2
    again = ClientDialog(store, data, None)
    again.code.setText("E7ZZ")
    again.name.setText("dup")
    assert "already a client E7ZZ" in again.save_findings(again.typed_findings())


def test_frame_view_turns_a_drag_into_fractions(qapp):
    from PySide6.QtCore import QPoint
    from PySide6.QtGui import QImage

    from src.app.client_dialogs import FrameView

    view = FrameView()
    view.resize(480, 360)                                     # its minimum size
    view.set_image(QImage(480, 360, QImage.Format.Format_RGB888))
    view._set_from(QPoint(24, 252), QPoint(456, 348))
    assert view.region == {"left": 0.05, "top": 0.7, "width": 0.9, "height": 0.267}
    parse_area(view.region)                                   # a valid region


def test_caption_area_dialog_on_a_real_video(qapp):
    video = FIXTURES / "sample_mistakes.mp4"
    if not video.is_file():
        pytest.skip("no sample video")
    from src.app.client_dialogs import load_frames

    frames = load_frames(video, count=3)
    assert frames and not frames[0][1].isNull()
