"""The Word version of the report (report.docx), read back with python-docx."""

import datetime as dt
import struct
import zlib
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm

from src.models import Caption, Issue
from src.report import builder
from src.report.builder import ReportInfo, build_report

WHEN = dt.datetime(2026, 10, 12, 9, 30, 0)
CAPTIONS = [
    Caption(1, 0.0, 2.0, "Welcome back."),
    Caption(2, 2.1, 4.1, "We need to to talk"),
    Caption(3, 4.2, 6.0, "Odd\x07 control char"),
    Caption(4, 6.1, 8.0, "Thanks."),
]
ISSUES = [
    Issue(2, 2.1, "repeated_word", "medium", "We need to to talk", 'Repeated word: "to to".',
          'Remove one: "We need to talk"', 0.8),
    Issue(2, 2.1, "extra_word", "high", "We need to to talk", 'Extra word "to" (repeated)', 'Audio: "We need to talk"', 1.0),
    Issue(3, 4.2, "spelling", "low", "Odd\x07 control char", "Possible misspelling", "", 0.5),
]


def make(tmp_path, formats=("md", "docx"), video=None, **kw):
    settings = {"output_dir": tmp_path / "reports", "formats": list(formats), "include_screenshots": True,
                "screenshot_at": "middle", "screenshot_width": 640, "write_csv": True, **kw}
    info = ReportInfo(srt=Path("clip.srt"), video=video, client="E3JB", rules_version="2026-10-12",
                      checks=["Spelling: ran"], generated=WHEN)
    return build_report(CAPTIONS, ISSUES, info, settings, ffmpeg="ffmpeg" if video else None)


def docx_of(written):
    [path] = [p for p in written if p.suffix == ".docx"]
    return path, Document(str(path))


def test_formats_from_settings(tmp_path):
    assert [p.name for p in make(tmp_path / "a")] == ["report.md", "report.docx", "issues.csv"]
    assert [p.name for p in make(tmp_path / "b", formats=["docx"])] == ["report.docx", "issues.csv"]
    assert [p.name for p in make(tmp_path / "c", formats=["md"], write_csv=False)] == ["report.md"]


def test_word_document_structure(tmp_path):
    _, doc = docx_of(make(tmp_path))
    headings = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith(("Title", "Heading"))]
    assert headings == [
        ("Title", "QC report: clip.srt"),
        ("Heading 1", "Summary"),
        ("Heading 1", "Flagged captions"),
        ("Heading 1", "Details"),
        ("Heading 2", "00:02.100, caption #2"),
        ("Heading 2", "00:04.200, caption #3"),
        ("Heading 1", "Appendix: all captions"),
    ]
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Client: E3JB" in text and "Rules version: 2026-10-12" in text
    assert "3 issues on 2 of 4 captions." in text
    assert doc.core_properties.title == "QC report: clip.srt"


def test_tables(tmp_path):
    _, doc = docx_of(make(tmp_path))
    summary, index, detail2, detail3, appendix = doc.tables
    assert [c.text for c in summary.rows[0].cells] == ["Issue type", "High", "Medium", "Low", "Total"]
    assert [c.text for c in summary.rows[-1].cells] == ["All", "1", "1", "1", "3"]
    assert [c.text for c in detail2.rows[1].cells][:2] == ["Extra word (audio)", "high"]   # worst first
    assert len(appendix.rows) == 1 + len(CAPTIONS)
    for t in doc.tables:                                   # header rows repeat on every page
        assert t.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None


def test_flagged_rows_are_bold_and_shaded(tmp_path):
    _, doc = docx_of(make(tmp_path))
    appendix = doc.tables[-1]
    flagged, plain = appendix.rows[2], appendix.rows[1]
    assert all(r.bold for r in flagged.cells[2].paragraphs[0].runs)
    assert flagged.cells[0]._tc.tcPr.find(qn("w:shd")).get(qn("w:fill")) == "FFF2CC"
    assert not any(r.bold for r in plain.cells[2].paragraphs[0].runs)


def test_links_point_to_bookmarks(tmp_path):
    _, doc = docx_of(make(tmp_path))
    body = doc.element.body
    bookmarks = {b.get(qn("w:name")) for b in body.iter(qn("w:bookmarkStart"))}
    anchors = {h.get(qn("w:anchor")) for h in body.iter(qn("w:hyperlink"))}
    assert bookmarks == {"caption_2", "caption_3"}
    assert anchors == bookmarks


def test_control_characters_are_removed(tmp_path):
    _, doc = docx_of(make(tmp_path))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Odd control char" in text and "\x07" not in text


def test_page_size(tmp_path):
    _, doc = docx_of(make(tmp_path / "a4"))
    assert abs(doc.sections[0].page_width - Cm(21)) < Cm(0.01)
    _, doc = docx_of(make(tmp_path / "letter", page_size="letter"))
    assert abs(doc.sections[0].page_width - Cm(21.59)) < Cm(0.01)


def tiny_png() -> bytes:
    """A valid 1x1 grey PNG, built with the standard library."""
    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))
    header = struct.pack(">IIBBBBB", 1, 1, 8, 0, 0, 0, 0)          # 1x1, 8-bit greyscale
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header)
            + chunk(b"IDAT", zlib.compress(b"\x00\x80")) + chunk(b"IEND", b""))


def test_screenshots_are_embedded(tmp_path, monkeypatch):
    png = tiny_png()

    def fake_grab(ffmpeg, video, seconds, out, width=640):
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_bytes(png)
        return True
    monkeypatch.setattr(builder, "grab_frame", fake_grab)
    _, doc = docx_of(make(tmp_path, video=Path("clip.mp4")))
    assert len(doc.inline_shapes) == 2
    assert "Frame at 00:03.100" in "\n".join(p.text for p in doc.paragraphs)


def test_settings_zoom_is_schema_valid(tmp_path):
    _, doc = docx_of(make(tmp_path))
    assert doc.settings.element.find(qn("w:zoom")).get(qn("w:percent")) == "100"


def test_screenshot_without_jfif_header_is_made_word_compatible(tmp_path):
    """ffmpeg 9 can write JPEGs with no JFIF header, which python-docx refuses."""
    from io import BytesIO

    from docx import Document
    from PIL import Image

    from src.report.screenshots import make_word_compatible

    buf = BytesIO()
    Image.new("RGB", (32, 18), (40, 60, 80)).save(buf, "JPEG")
    data = buf.getvalue()
    assert data[2:4] == b"\xff\xe0"                       # APP0 (JFIF) segment
    app0_len = int.from_bytes(data[4:6], "big")
    bare = tmp_path / "bare.jpg"
    bare.write_bytes(data[:2] + data[4 + app0_len:])      # like ffmpeg 9: no JFIF header
    assert make_word_compatible(bare)
    Document().add_picture(str(bare))                    # would raise UnrecognizedImageError before
