"""The engine entry point run_qc: progress, cancelling, errors, options, and the Word report.
Uses an SRT or a stubbed OCR, so no real video work happens here."""

import threading
from pathlib import Path

import pytest
from docx import Document

from src.engine import (STAGES, QCError, QCOptions, build_docx_report, guess_client, known_clients, run_qc,
                        write_report)
from src.models import Caption, SpokenWord, WRONG_WORD
from src.utils.progress import Cancelled, Stage

FIXTURES = Path(__file__).resolve().parent / "fixtures"
SRT = FIXTURES / "sample_mistakes.srt"
STAGE_NAMES = [name for name, _ in STAGES]


def no_llm(**kw):
    return QCOptions(srt_path=SRT, use_llm=False, **kw)


def test_run_qc_returns_the_issue_list_sorted_by_time():
    result = run_qc(None, "E8D", options=no_llm())
    assert isinstance(result, list) and len(result) > 0
    assert result == sorted(result, key=lambda i: (i.timestamp, i.caption_id))
    assert result.issues == list(result)
    assert len(result.captions) == 20 and result.client["client"] == "E8D"
    assert any(line.startswith("Spelling and repeated words: ran") for line in result.checks)


def test_progress_goes_through_the_stages_in_order_and_ends_at_one():
    calls = []
    run_qc(None, None, lambda stage, f: calls.append((stage, f)), options=no_llm())
    fractions = [f for _, f in calls]
    assert fractions == sorted(fractions) and fractions[-1] == pytest.approx(1.0)
    seen = list(dict.fromkeys(stage for stage, _ in calls))
    assert seen == STAGE_NAMES                              # stages that don't apply still pass by
    assert all(0.0 <= f <= 1.0 for f in fractions)


def test_cancel_before_starting():
    stop = threading.Event()
    stop.set()
    with pytest.raises(Cancelled):
        run_qc(None, None, cancel_event=stop, options=no_llm())


def test_cancel_during_the_run():
    stop = threading.Event()

    def progress(stage, fraction):
        if stage == "Checking captions":
            stop.set()

    with pytest.raises(Cancelled, match="Checking captions"):
        run_qc(None, None, progress, stop, options=no_llm())


def test_cancel_reaches_the_ocr_loop(tmp_path, monkeypatch):
    """The OCR reports progress per frame, so a cancel stops it mid-video."""
    video = tmp_path / "clip.mp4"
    video.write_bytes(b"fake")
    stop = threading.Event()
    frames_read = []

    def fake_ocr(path, settings, cache_dir, progress: Stage = None, **kw):
        for n in range(100):
            if n == 5:
                stop.set()
            progress(n / 100)
            frames_read.append(n)
        return []

    monkeypatch.setattr("src.engine.read_burned_in_captions", fake_ocr)
    with pytest.raises(Cancelled, match="Reading captions"):
        run_qc(video, None, cancel_event=stop, options=QCOptions(use_llm=False))
    assert len(frames_read) == 5


def test_errors_are_readable():
    with pytest.raises(QCError, match="File not found"):
        run_qc(Path("no such video.mp4"), None)
    with pytest.raises(QCError, match="Give a video"):
        run_qc(None, None)
    with pytest.raises(QCError, match="SRT file not found"):
        run_qc(None, None, options=QCOptions(srt_path=Path("missing.srt")))


def test_unreadable_video_is_a_qc_error(tmp_path, monkeypatch):
    video = tmp_path / "broken.mp4"
    video.write_bytes(b"not a video")

    def broken(*a, **kw):
        raise RuntimeError("ffmpeg could not extract audio from broken.mp4")

    monkeypatch.setattr("src.engine.read_burned_in_captions", lambda *a, **kw: [Caption(1, 0, 1, "Hi there")])
    monkeypatch.setattr("src.engine.get_transcript", broken)
    with pytest.raises(QCError, match="ffmpeg could not extract audio"):
        run_qc(video, None, options=QCOptions(use_llm=False))


def test_checks_can_be_switched_off_per_run():
    every = run_qc(None, None, options=no_llm())
    fewer = run_qc(None, None, options=no_llm(disabled_checks=frozenset({"spelling", "too_short"})))
    assert {"spelling", "too_short"} & {i.type for i in every}
    assert not {"spelling", "too_short", "phrase_break", "repeated_word"} & {i.type for i in fewer}
    assert "Spelling and repeated words: switched off for this run" in fewer.checks


def test_docx_report_path_has_the_video_name_and_date(tmp_path):
    result = run_qc(None, "E8D", options=no_llm())
    path = build_docx_report(result, output_dir=tmp_path)
    assert path.is_file() and path.suffix == ".docx"
    assert path.name.startswith("sample_mistakes_QC_20") and path.parent.parent == tmp_path
    doc = Document(str(path))
    assert [h.text for h in doc.paragraphs if h.style.name == "Heading 1"][:2] == ["Summary", "Issues"]
    assert len(doc.tables[1].rows) == 1 + len(result)


def test_paths_with_spaces_and_non_english_characters(tmp_path):
    folder = tmp_path / "Vidéos à vérifier" / "客户 E8D"
    folder.mkdir(parents=True)
    srt = folder / "clip été – 01.srt"
    srt.write_bytes(SRT.read_bytes())
    result = run_qc(None, "E8D", options=QCOptions(srt_path=srt, use_llm=False))
    written = write_report(result, formats=["md", "docx"], output_dir=folder / "rapports")
    assert all(p.is_file() for p in written)
    assert any(p.name.startswith("clip été – 01_QC_") for p in written)


def test_guess_client_from_the_file_name():
    known = ["E2CI", "E2OCA", "E3RK", "E8C", "E8D"]
    assert guess_client(Path("E8D_reel_01.mp4"), known) == "E8D"
    assert guess_client(Path("8d - reel final.mov"), known) == "E8D"
    assert guess_client(Path("Rev2 E2OCA interview.mp4"), known) == "E2OCA"
    assert guess_client(Path("E8DX reel.mp4"), known) is None              # not a whole code
    assert guess_client(Path("del.mp4"), known) is None
    assert "E8D" in known_clients()


# --- the speech skip rule, through the engine -----------------------------------------------

SPEECH_TYPES = {"wrong_word", "missing_word", "extra_word", "caption_early", "caption_late",
                "caption_sync", "uncaptioned_speech"}


def spoken(text: str, gap: float = 0.5, length: float = 0.35, probability: float = 0.9) -> list[SpokenWord]:
    return [SpokenWord(word, i * gap, i * gap + length, probability) for i, word in enumerate(text.split())]


def stub_speech(tmp_path, monkeypatch, words, captions):
    """A fake video whose captions and transcript are what the test says."""
    video = tmp_path / "clip.mp4"
    video.write_bytes(b"fake")
    monkeypatch.setattr("src.engine.read_burned_in_captions", lambda *a, **kw: captions)
    monkeypatch.setattr("src.engine.get_transcript", lambda *a, **kw: words)
    return video


def test_one_word_transcript_skips_the_speech_checks_but_not_the_others(tmp_path, monkeypatch):
    video = stub_speech(tmp_path, monkeypatch, [SpokenWord("Yeah", 0.5, 0.7, 0.9)],
                        [Caption(1, 0.0, 2.0, "i went to teh store")])
    result = run_qc(video, None, options=QCOptions(use_llm=False))
    assert not SPEECH_TYPES & {i.type for i in result}
    assert any(note.startswith("Speech comparison skipped:") for note in result.notes)
    assert any("Captions vs audio (Whisper): skipped (only 1 word" in line for line in result.checks)
    assert any(line.startswith("Spelling and repeated words: ran") for line in result.checks)   # others still ran


def test_a_normal_transcript_runs_the_speech_checks(tmp_path, monkeypatch):
    caption = "i went to teh store and the market and home"
    video = stub_speech(tmp_path, monkeypatch, spoken("i went to the store and the market and home"),
                        [Caption(1, 0.0, 6.0, caption)])
    result = run_qc(video, None, options=QCOptions(use_llm=False))
    assert result.notes == []                                                     # nothing skipped
    assert any("Captions vs audio (Whisper):" in line and "10 words heard" in line for line in result.checks)
    assert WRONG_WORD in {i.type for i in result}                                 # the comparison really ran


def test_a_video_with_no_audio_track_skips_the_comparison(tmp_path, monkeypatch):
    video = tmp_path / "silent.mp4"
    video.write_bytes(b"fake")
    monkeypatch.setattr("src.engine.read_burned_in_captions", lambda *a, **kw: [Caption(1, 0.0, 1.0, "Hi there")])
    monkeypatch.setattr("src.engine.has_audio", lambda *a, **kw: False)
    asked = []
    monkeypatch.setattr("src.engine.get_transcript", lambda *a, **kw: asked.append(1))
    result = run_qc(video, None, options=QCOptions(use_llm=False))
    assert asked == []                                                            # never transcribed
    assert result.notes and "no audio track" in result.notes[0]


def test_low_probability_words_skip_the_comparison(tmp_path, monkeypatch):
    video = stub_speech(tmp_path, monkeypatch, spoken("word " * 40, probability=0.2),
                        [Caption(1, 0.0, 5.0, "some caption here for the test")])
    monkeypatch.setattr("src.engine.media_duration", lambda *a, **kw: 27.0)        # enough coverage, unsure words
    result = run_qc(video, None, options=QCOptions(use_llm=False))
    assert result.notes and "guesses over music" in result.notes[0]


def test_speech_limits_can_be_overridden_for_a_client(tmp_path, monkeypatch):
    caption = "one two three four five six seven eight nine ten eleven twelve"
    video = stub_speech(tmp_path, monkeypatch, spoken(caption), [Caption(1, 0.0, 6.0, caption)])
    monkeypatch.setattr("src.engine.media_duration", lambda *a, **kw: 27.0)
    client = {"client": "E9Z", "overrides": {"transcript": {"min_words": 50}}}
    monkeypatch.setattr("src.engine.load_client",
                        lambda config, code: client if code else {"client": None, "overrides": {}})
    assert run_qc(video, None, options=QCOptions(use_llm=False)).notes == []       # 12 >= 10: runs by default
    strict = run_qc(video, "E9Z", options=QCOptions(use_llm=False))
    assert strict.notes and "only 12 words" in strict.notes[0]                     # 12 < 50: the client's limit
