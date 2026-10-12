"""The QC report as a Word document (report.docx), with the same content as report.md.

Uses python-docx. Built-in heading styles (so Word's Navigation Pane works), real tables with
fixed column widths and header rows that repeat on every page, screenshots embedded in the file,
links from the index and appendix to each caption's section, and page numbers in the footer.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from src.models import SEVERITIES
from src.report.builder import ReportData, label, worst
from src.utils.timecode import format_timestamp

PAGE_SIZES = {"a4": (21.0, 29.7), "letter": (21.59, 27.94)}     # cm
MARGIN_CM = 2.0
SEVERITY_COLORS = {"high": "C00000", "medium": "C55A11", "low": "7F7F7F"}
HEADER_FILL = "D9E2F3"       # light blue header rows
FLAGGED_FILL = "FFF2CC"      # light yellow for flagged captions in the appendix
LINK_COLOR = "0563C1"
_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def clean(text) -> str:
    """Text Word can store: one line, no control characters (they would make the file invalid)."""
    return _CONTROL_CHARS.sub("", " ".join(str(text).split()))


# --- small XML helpers python-docx doesn't offer -------------------------------------

def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")           # "clear" + fill; "solid" would paint the cell black
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:tblHeader")
    flag.set(qn("w:val"), "true")
    tr_pr.append(flag)


def bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    p = paragraph._p
    p_pr = p.find(qn("w:pPr"))
    if p_pr is not None:
        p_pr.addnext(start)
    else:
        p.insert(0, start)
    p.append(end)


def internal_link(paragraph, text: str, anchor: str, bold: bool = False) -> None:
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), anchor)
    link.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), LINK_COLOR)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    if bold:
        r_pr.append(OxmlElement("w:b"))
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = clean(text)
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def page_number_field(paragraph) -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    run.append(t)
    field.append(run)
    paragraph._p.append(field)


def anchor_name(cid: int) -> str:
    return f"caption_{cid}"


# --- tables ---------------------------------------------------------------------------

def table(doc, headers: list[str], widths_cm: list[float], right_align: set[int] = frozenset()):
    """A grid table with a shaded, repeating header row and fixed column widths."""
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, (cell, text) in enumerate(zip(t.rows[0].cells, headers)):
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        if i in right_align:
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        shade(cell, HEADER_FILL)
    repeat_header(t.rows[0])
    t._widths = widths_cm
    t._right = set(right_align)
    set_widths(t)
    return t


def set_widths(t) -> None:
    for row in t.rows:
        for cell, width in zip(row.cells, t._widths):
            cell.width = Cm(width)
    for column, width in zip(t.columns, t._widths):
        column.width = Cm(width)


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def add_row(t, values: list, bold: bool = False, fill: str | None = None, colors: dict | None = None):
    row = t.add_row()
    keep_row_together(row)
    for i, (cell, value) in enumerate(zip(row.cells, values)):
        cell.width = Cm(t._widths[i])
        paragraph = cell.paragraphs[0]
        if i in t._right:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if callable(value):
            value(paragraph)                         # e.g. a link
        else:
            run = paragraph.add_run(clean(value))
            run.bold = bold
            if colors and i in colors:
                run.font.color.rgb = RGBColor.from_string(colors[i])
                run.bold = True
        if fill:
            shade(cell, fill)
    return row


def link_cell(text: str, cid: int, bold: bool = False):
    return lambda paragraph: internal_link(paragraph, text, anchor_name(cid), bold)


# --- the document -----------------------------------------------------------------------

def write_docx(data: ReportData) -> Path:
    settings, info = data.settings, data.info
    doc = Document()
    width_cm, height_cm = PAGE_SIZES.get(str(settings.get("page_size", "a4")).lower(), PAGE_SIZES["a4"])
    content_cm = width_cm - 2 * MARGIN_CM
    for section in doc.sections:
        section.page_width, section.page_height = Cm(width_cm), Cm(height_cm)
        section.left_margin = section.right_margin = Cm(MARGIN_CM)
        section.top_margin = section.bottom_margin = Cm(MARGIN_CM)
    # python-docx's built-in template leaves out the zoom percentage the Word schema requires.
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)

    title = (info.video or info.srt).name
    doc.core_properties.title = f"QC report: {title}"
    doc.core_properties.author = "Video QC Bot"

    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = f"QC report: {clean(title)}    |    page "
    page_number_field(footer)
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Header
    doc.add_heading(f"QC report: {clean(title)}", level=0)
    facts = [
        ("Video", info.video.name if info.video else "none (caption text checks only)"),
        ("Captions", f"{info.srt.name}, {len(data.captions)} captions"),
        ("Client", info.client or "default rules"),
        ("Generated", f"{info.generated:%Y-%m-%d %H:%M}"),
        ("Rules version", info.rules_version or "unknown"),
    ]
    for name, value in facts:
        p = doc.add_paragraph()
        p.add_run(f"{name}: ").bold = True
        p.add_run(clean(value))
        p.paragraph_format.space_after = Pt(0)
    if info.checks:
        p = doc.add_paragraph()
        p.add_run("Checks:").bold = True
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(0)
        for line in info.checks:
            doc.add_paragraph(clean(line), style="List Bullet")
    n = len(data.issues)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.add_run(f"{n} issue{'s' if n != 1 else ''} on {len(data.flagged)} of {len(data.captions)} captions.").bold = True

    # Summary
    doc.add_heading("Summary", level=1)
    if not data.issues:
        doc.add_paragraph("No issues found.")
    else:
        counts = data.counts()
        rest = (content_cm - 6.0) / 4
        t = table(doc, ["Issue type", "High", "Medium", "Low", "Total"], [6.0, rest, rest, rest, rest],
                  right_align={1, 2, 3, 4})
        for kind in data.types():
            row = [counts[(kind, s)] for s in SEVERITIES]
            add_row(t, [label(kind)] + [str(v or "") for v in row] + [str(sum(row))])
        totals = [sum(counts[(k, s)] for k in data.types()) for s in SEVERITIES]
        add_row(t, ["All"] + [str(v) for v in totals] + [str(sum(totals))], bold=True)
        doc.add_paragraph("Severity: high is almost always a real mistake, medium usually is, low is worth a "
                          "glance (often a name to add to the glossary).").paragraph_format.space_before = Pt(6)

    # Index of flagged captions
    if data.flagged:
        doc.add_heading("Flagged captions", level=1)
        t = table(doc, ["Time", "Caption", "Problems", "Worst"], [2.2, content_cm - 2.2 - 4.8 - 1.8, 4.8, 1.8])
        for cid in data.flagged:
            group = data.by_caption[cid]
            problems = ", ".join(dict.fromkeys(label(i.type) for i in group))
            text = data.caption_text(cid)
            short = text if len(text) <= 70 else text[:67] + "..."
            sev = worst(group)
            add_row(t, [link_cell(format_timestamp(group[0].timestamp), cid), f"#{cid} {short}", problems, sev],
                    colors={3: SEVERITY_COLORS[sev]})

    # Details, one section per flagged caption
    if data.flagged:
        doc.add_heading("Details", level=1)
        for n_section, cid in enumerate(data.flagged, start=1):
            group = data.by_caption[cid]
            heading = doc.add_heading(f"{format_timestamp(group[0].timestamp)}, caption #{cid}", level=2)
            bookmark(heading, anchor_name(cid), n_section)
            # Keep each caption's heading, text, screenshot and table on the same page.
            on_screen = doc.add_paragraph(f"On screen {data.on_screen(cid)}")
            on_screen.paragraph_format.space_after = Pt(2)
            on_screen.paragraph_format.keep_with_next = True
            quote = doc.add_paragraph(style="Quote")
            quote.add_run(clean(data.caption_text(cid)))
            quote.paragraph_format.keep_with_next = True
            if cid in data.shots:
                image, at = data.shots[cid]
                doc.add_picture(str(data.folder / image), width=Cm(min(12.0, content_cm)))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                doc.paragraphs[-1].paragraph_format.keep_with_next = True
                note = doc.add_paragraph()
                note.alignment = WD_ALIGN_PARAGRAPH.CENTER
                note.add_run(f"Frame at {format_timestamp(at)}").italic = True
                note.paragraph_format.keep_with_next = True
            elif data.wanted_shots:
                missing = doc.add_paragraph()
                missing.add_run("(no screenshot: ffmpeg could not get a frame at this time)").italic = True
                missing.paragraph_format.keep_with_next = True
            fixed = 3.2 + 1.8 + 2.3
            t = table(doc, ["Problem", "Severity", "What's wrong", "Suggested fix", "Confidence"],
                      [3.2, 1.8, (content_cm - fixed) * 0.55, (content_cm - fixed) * 0.45, 2.3], right_align={4})
            for issue in group:
                add_row(t, [label(issue.type), issue.severity, issue.message, issue.suggestion,
                            f"{issue.confidence:.2f}"], colors={1: SEVERITY_COLORS[issue.severity]})

    # Appendix
    doc.add_heading("Appendix: all captions", level=1)
    doc.add_paragraph("Flagged captions are in bold on a yellow background.")
    t = table(doc, ["Time", "#", "Caption", "Issues"], [2.2, 1.2, content_cm - 2.2 - 1.2 - 4.6, 4.6],
              right_align={1})
    for cap in data.captions:
        group = data.by_caption.get(cap.id, [])
        if group:
            found = ", ".join(dict.fromkeys(label(i.type) for i in group))
            add_row(t, [format_timestamp(cap.start), link_cell(str(cap.id), cap.id, bold=True),
                        cap.flat_text, found], bold=True, fill=FLAGGED_FILL)
        else:
            add_row(t, [format_timestamp(cap.start), str(cap.id), cap.flat_text, ""])

    path = data.folder / "report.docx"
    doc.save(str(path))
    return path
